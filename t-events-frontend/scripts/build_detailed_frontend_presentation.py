from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = Path(r"C:\Users\Office\Downloads\Telegram Desktop\Чанышева_Презентация_Т-евентс.pptx")
FALLBACK = Path(r"C:\Users\Office\Downloads\02_HSE_Presentation_Shablon.pptx")
OUT = ROOT / "docs" / "t_events_frontend_detailed_defense.pptx"

NAVY = RGBColor(16, 45, 105)
LIME = RGBColor(220, 255, 5)
LAV = RGBColor(223, 199, 242)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(16, 16, 16)
GRAY = RGBColor(243, 245, 248)
PALE = RGBColor(248, 249, 252)

AUTHOR = "Автор: ____________________\nОП «Программная инженерия»\nГруппа: __________"
FOOTER_TITLE = "«Клиентская часть приложения T-Events для участия в образовательных активностях c элементами геймификации»"
CITY_YEAR = "Москва, 2026"


SLIDES = [
    {
        "kind": "cover",
        "title": "«Клиентская часть приложения T-Events для участия в образовательных активностях c элементами геймификации»",
        "subtitle": "Frontend-часть программного проекта",
    },
    {
        "title": "Описание предметной области",
        "blocks": [
            ("Контекст", [
                "Т-Образование проводит стендовые активности на Днях карьеры, днях открытых дверей и партнерских мероприятиях.",
                "Участники хотят быстро понять доступные активности, пройти задания и получить фирменную продукцию.",
            ]),
            ("Проблема", [
                "Ручной учет прогресса и выдачи мерча плохо масштабируется.",
                "Сложно учитывать направление участника, баллы, доступность награды и факт выдачи.",
            ]),
            ("Frontend-задача", [
                "Собрать путь участника, стендиста и администратора в едином интерфейсе.",
                "Сделать клиентский слой устойчивым к ошибкам API, авторизации и смене состояния мероприятия.",
            ]),
        ],
    },
    {
        "title": "Цель и задачи",
        "left": ("Цель курсовой работы", [
            "Реализовать клиентскую часть приложения T-Events, обеспечивающую участие в мероприятиях, прохождение игровых активностей, получение QR-наград и управление мероприятиями.",
        ]),
        "right": ("Задачи", [
            "Спроектировать frontend-архитектуру по ролям и пользовательским сценариям.",
            "Реализовать интеграцию с API авторизации, мероприятий, игр, наград и администрирования.",
            "Реализовать игровые сессии для question_answer и quiz.",
            "Реализовать QR/redeem-сценарий выдачи призов.",
            "Добавить runtime-проверку критичных API-контрактов.",
        ]),
    },
    {
        "title": "Анализ аналогов",
        "blocks": [
            ("Прямой аналог", [
                "T-education_bot: закрывает часть сценариев, но не дает полноценного веб-интерфейса с ролями, направлениями, QR-выдачей и администрированием.",
            ]),
            ("Косвенные аналоги", [
                "Stepik, Quizizz, Duolingo: сильны в обучении и игровых заданиях, но не решают задачу стендового мероприятия и физической выдачи мерча.",
            ]),
            ("Вывод", [
                "T-Events является узкоспециализированным решением: оно объединяет участие, геймификацию, прогресс, право на приз и подтвержденную выдачу.",
            ]),
        ],
    },
    {
        "title": "Анализ предметной области. Пользовательские роли",
        "roles": [
            ("Участник", ["Каталог мероприятий", "Выбор направления", "Прохождение игр", "QR-награда"]),
            ("Стендист", ["Сканирование QR", "Ручной redeem-код", "Preview выдачи", "Журнал призов"]),
            ("Администратор", ["Создание мероприятия", "Направления и игры", "Пороги наград", "Публикация/audit"]),
        ],
        "note": "Ключевая сложность frontend-части: три роли работают с общими сущностями, но имеют разные права, интерфейсы и критичные действия.",
    },
    {
        "title": "Пользовательский путь участника",
        "flow": ["Каталог", "Мероприятие", "Направление", "Игры", "Сессия", "Прогресс", "QR-награда"],
        "details": [
            "Состояние участия сохраняет выбранное мероприятие и направление.",
            "Игровая сессия может запускаться или продолжаться с backend.",
            "После ответа обновляется прогресс игры и summary направления.",
            "Доступность награды определяется порогами малой/большой награды.",
        ],
    },
    {
        "title": "Пользовательский путь стендиста",
        "flow": ["Вход", "Сканер/код", "Preview", "Confirm", "Redemption", "Inventory"],
        "details": [
            "QR содержит signed token, но предусмотрен ручной redeem-код.",
            "Выдача разделена на два шага: предварительная проверка и подтверждение.",
            "После подтверждения создается запись выдачи, доступная в журнале.",
            "Такой сценарий снижает риск ошибочной или повторной выдачи.",
        ],
    },
    {
        "title": "Пользовательский путь администратора",
        "flow": ["События", "Детали", "Расписание", "Направления", "Игры", "Readiness", "Публикация"],
        "details": [
            "Мероприятие нельзя считать готовым без расписания, порогов наград, направлений и игр.",
            "Frontend отображает readiness-модель и группирует проблемы по разделам.",
            "Администратор видит, какие настройки мешают публикации.",
            "Audit/export позволяет анализировать административные действия.",
        ],
    },
    {
        "title": "Выбор технологий для реализации",
        "left": ("Инструментарий frontend", [
            "Next.js 16 App Router, React 19, TypeScript 5",
            "Tailwind CSS 4, собственные UI-примитивы",
            "TanStack React Query 5 для server state",
            "Zustand для локального состояния участия",
            "zod для runtime API-контрактов",
            "qrcode.react и @zxing/browser для QR-сценариев",
        ]),
        "right": ("Обоснование выбора", [
            "App Router удобен для route-level состояний loading/error/not-found.",
            "React Query отделяет backend state от локального UI state.",
            "zod закрывает проблему недоверенных backend-ответов.",
            "TypeScript strict снижает риск ошибок при росте DTO и сценариев.",
        ]),
    },
    {
        "title": "Архитектура клиентского приложения",
        "columns": [
            ("src/app", ["Маршруты", "page.tsx как route shell", "loading/error/not-found", "Client-контейнеры"]),
            ("src/features", ["auth/events/reward/admin", "API + contracts", "hooks", "presentation helpers"]),
            ("src/lib", ["api/client", "api/types", "routes", "queryKeys", "tokenStore"]),
            ("src/components", ["UI primitives", "layout", "route states", "events visuals"]),
        ],
        "note": "Архитектурный принцип: route-слой связывает сценарий, feature-слой содержит предметную логику, lib-слой содержит инфраструктуру.",
    },
    {
        "title": "Архитектура API-интеграции и состояния",
        "flow": ["UI", "Feature hook", "React Query", "API client", "zod contract", "Backend"],
        "details": [
            "API-клиент добавляет Authorization, JSON headers, CSRF token и credentials.",
            "При 401 выполняется refresh access token и повтор исходного запроса.",
            "React Query кэширует server state и инвалидирует данные после мутаций.",
            "Zustand используется только для локального сценария участия, а не как копия backend.",
        ],
    },
    {
        "title": "Сложная часть 1. Игровая сессия",
        "left": ("Трудности", [
            "Сессия не является одной страницей: есть статус, прогресс, навигация, текущий вопрос и история ответов.",
            "Нужно поддержать два движка: question_answer и quiz.",
            "После ответа меняется не только вопрос, но и прогресс направления.",
        ]),
        "right": ("Решение", [
            "DTO-модель SessionState / StartOrResume / SubmitAnswer.",
            "discriminated union по полю engine.",
            "Нормализация session DTO в runtime state.",
            "Feature helpers для выбора текущего вопроса и отображения результата.",
        ]),
    },
    {
        "title": "Сложная часть 2. QR-награды и выдача",
        "left": ("Трудности", [
            "Награда зависит от прогресса направления, а не только от отдельной игры.",
            "QR должен быть связан с мероприятием, направлением, типом награды и пользователем.",
            "Нужно поддержать камеру и ручной код без потери надежности сценария.",
        ]),
        "right": ("Решение", [
            "RewardEligibility: locked/small_unlocked/big_unlocked/redeemed.",
            "RewardQr: signed_token, redeem_code, expires_at, status.",
            "Двухшаговая выдача: preview -> confirm.",
            "Redemption inventory для контроля выданных призов.",
        ]),
    },
    {
        "title": "Сложная часть 3. Авторизация и роли",
        "left": ("Трудности", [
            "Разные сценарии доступны разным ролям: participant, stander, admin.",
            "Сессия может истечь во время работы с приложением.",
            "Несколько вкладок должны корректно реагировать на login/logout.",
        ]),
        "right": ("Решение", [
            "AuthProvider с явной моделью статуса авторизации.",
            "Access token в памяти, refresh token через cookie.",
            "Refresh coalescing: один refresh для параллельных 401.",
            "BroadcastChannel и storage events для синхронизации вкладок.",
        ]),
    },
    {
        "title": "Сложная часть 4. Runtime API-контракты",
        "left": ("Проблема", [
            "TypeScript проверяет код, но не гарантирует корректность данных от backend.",
            "Ошибка структуры ответа в игре, награде или публикации может привести к неверному UI-состоянию.",
        ]),
        "right": ("Реализация", [
            "authContracts: login/register/me.",
            "eventContracts: event, directions, games, session, submit answer.",
            "rewardContracts: eligibility, QR, preview, redemption, inventory.",
            "adminContracts: event settings, readiness, games, publish check.",
        ]),
    },
    {
        "title": "Основные реализованные модули",
        "table": [
            ("Auth", "Email/password, VK callback, refresh, роли, logout sync"),
            ("Events", "Каталог, карточка, направления, игры, сессии, leaderboard"),
            ("Reward", "Статус награды, QR, redeem-код, сканер, журнал выдач"),
            ("Admin", "События, настройки, направления, игры, публикация, архив, audit"),
            ("UI/UX", "Route states, error/empty/loading, формы, доступные примитивы"),
        ],
    },
    {
        "title": "Результаты проделанной работы",
        "blocks": [
            ("Результат 1", ["Сформирована клиентская архитектура, соответствующая ролям и бизнес-процессам T-Events."]),
            ("Результат 2", ["Реализован полный путь участника: мероприятие -> направление -> игра -> прогресс -> QR-награда."]),
            ("Результат 3", ["Реализован операционный путь стендиста: QR/redeem -> preview -> confirm -> inventory."]),
            ("Результат 4", ["Реализован административный контур: настройки, readiness, публикация, архивирование и audit/export."]),
        ],
    },
    {
        "title": "Демонстрация работы продукта",
        "demo": [
            ("1", "Участник выбирает мероприятие и направление"),
            ("2", "Проходит игровую сессию и получает прогресс"),
            ("3", "Открывает QR-награду"),
            ("4", "Стендист проверяет redeem-код и подтверждает выдачу"),
            ("5", "Администратор видит настройки, готовность публикации и журнал"),
        ],
        "note": "Демонстрация показывает сквозной процесс, ради которого строилась frontend-архитектура.",
    },
    {
        "title": "Список источников",
        "sources": [
            "Next.js App Router Documentation. URL: https://nextjs.org/docs/app",
            "React Documentation. URL: https://react.dev/",
            "TanStack Query Documentation. URL: https://tanstack.com/query",
            "Zod Documentation. URL: https://zod.dev/",
            "ZXing Browser Documentation. URL: https://github.com/zxing-js/browser",
            "WCAG 2.2. URL: https://www.w3.org/TR/WCAG22/",
        ],
    },
]


def clear_slides(prs):
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)


def blank_layout(prs):
    for layout in prs.slide_layouts:
        if len(layout.placeholders) == 0:
            return layout
    return prs.slide_layouts[-1]


def add_rect(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    return shape


def add_text(slide, x, y, w, h, text, size=14, color=BLACK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_footer(slide, n):
    add_text(slide, 0.62, 6.18, 3.4, 0.45, AUTHOR, 6.5, NAVY)
    add_text(slide, 0.62, 6.82, 1.4, 0.18, CITY_YEAR, 6.5, NAVY)
    add_text(slide, 3.0, 6.8, 7.7, 0.3, FOOTER_TITLE, 6.5, NAVY)
    add_text(slide, 12.15, 6.82, 0.55, 0.25, str(n), 8, NAVY, True, PP_ALIGN.RIGHT)


def add_title(slide, title):
    add_text(slide, 0.62, 0.45, 11.5, 0.55, title, 25, NAVY, True)


def add_bullet_list(slide, x, y, w, items, size=13.2, gap=0.38):
    for item in items:
        add_text(slide, x, y, w, 0.32, "• " + item, size, BLACK)
        y += gap
    return y


def add_blocks(slide, blocks):
    y = 1.2
    for heading, items in blocks:
        add_rect(slide, 0.65, y, 11.85, 1.05, PALE, LAV)
        add_text(slide, 0.85, y + 0.1, 2.0, 0.25, heading, 13.5, NAVY, True)
        add_bullet_list(slide, 2.8, y + 0.1, 9.3, items, 12.4, 0.32)
        y += 1.18


def add_two_columns(slide, left, right):
    for x, data in [(0.72, left), (6.75, right)]:
        heading, items = data
        add_rect(slide, x, 1.25, 5.55, 4.6, PALE, LAV)
        add_text(slide, x + 0.2, 1.42, 5.1, 0.32, heading, 15.5, NAVY, True)
        add_bullet_list(slide, x + 0.25, 1.9, 5.05, items, 12.8, 0.43)


def add_roles(slide, roles, note):
    x = 0.72
    for heading, items in roles:
        add_rect(slide, x, 1.3, 3.75, 3.45, PALE, LAV)
        add_text(slide, x + 0.18, 1.5, 3.35, 0.35, heading, 17, NAVY, True, PP_ALIGN.CENTER)
        add_bullet_list(slide, x + 0.28, 2.1, 3.15, items, 12.7, 0.43)
        x += 4.05
    add_rect(slide, 0.72, 5.15, 11.85, 0.68, NAVY)
    add_text(slide, 0.92, 5.3, 11.45, 0.3, note, 12.2, WHITE)


def add_flow(slide, flow, details):
    y = 1.45
    x = 0.7
    step_w = 1.45 if len(flow) > 6 else 1.65
    gap = 0.25
    for i, label in enumerate(flow):
        add_rect(slide, x, y, step_w, 0.58, LAV if i % 2 == 0 else PALE, NAVY)
        add_text(slide, x + 0.04, y + 0.16, step_w - 0.08, 0.18, label, 10.5, NAVY, True, PP_ALIGN.CENTER)
        if i < len(flow) - 1:
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + step_w), Inches(y + 0.29), Inches(x + step_w + gap), Inches(y + 0.29))
            conn.line.color.rgb = NAVY
            conn.line.width = Pt(1.25)
        x += step_w + gap
    add_rect(slide, 0.75, 2.65, 11.65, 2.65, PALE, LAV)
    add_bullet_list(slide, 1.0, 2.95, 10.95, details, 13.2, 0.5)


def add_columns(slide, columns, note):
    x = 0.55
    width = 3.05
    for heading, items in columns:
        add_rect(slide, x, 1.25, width, 3.9, PALE, LAV)
        add_text(slide, x + 0.12, 1.42, width - 0.24, 0.28, heading, 13.8, NAVY, True, PP_ALIGN.CENTER)
        add_bullet_list(slide, x + 0.18, 1.9, width - 0.3, items, 11.3, 0.43)
        x += 3.18
    add_rect(slide, 0.65, 5.45, 11.95, 0.46, NAVY)
    add_text(slide, 0.85, 5.55, 11.55, 0.18, note, 10.5, WHITE)


def add_table(slide, rows):
    y = 1.2
    add_rect(slide, 0.75, y, 11.75, 0.42, NAVY)
    add_text(slide, 0.95, y + 0.08, 2.2, 0.16, "Модуль", 10.5, WHITE, True)
    add_text(slide, 3.1, y + 0.08, 9.0, 0.16, "Реализованная функциональность", 10.5, WHITE, True)
    y += 0.48
    for name, desc in rows:
        add_rect(slide, 0.75, y, 2.15, 0.55, PALE, LAV)
        add_rect(slide, 2.9, y, 9.6, 0.55, WHITE, LAV)
        add_text(slide, 0.95, y + 0.13, 1.7, 0.16, name, 10.8, NAVY, True)
        add_text(slide, 3.1, y + 0.13, 9.0, 0.16, desc, 10.8, BLACK)
        y += 0.62


def add_demo(slide, demo, note):
    x = 0.9
    for num, text in demo:
        add_rect(slide, x, 1.45, 2.1, 2.7, PALE, LAV)
        add_text(slide, x + 0.68, 1.75, 0.7, 0.55, num, 30, NAVY, True, PP_ALIGN.CENTER)
        add_text(slide, x + 0.22, 2.62, 1.65, 0.8, text, 11.5, BLACK, False, PP_ALIGN.CENTER)
        x += 2.35
    add_rect(slide, 0.75, 5.25, 11.75, 0.58, NAVY)
    add_text(slide, 0.95, 5.38, 11.3, 0.23, note, 12, WHITE)


def add_sources(slide, sources):
    add_bullet_list(slide, 0.85, 1.25, 11.1, sources, 12.2, 0.47)


def build():
    template = REFERENCE if REFERENCE.exists() else FALLBACK
    prs = Presentation(str(template)) if template.exists() else Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    clear_slides(prs)
    layout = blank_layout(prs)

    for i, data in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(layout)
        if data.get("kind") == "cover":
            add_rect(slide, 0, 0, 13.333, 7.5, NAVY)
            add_rect(slide, 0, 0, 13.333, 0.25, LIME)
            add_text(slide, 0.75, 1.2, 10.2, 1.6, data["title"], 27, WHITE, True)
            add_text(slide, 0.78, 3.18, 7.4, 0.42, data["subtitle"], 17, LIME, True)
            add_text(slide, 0.78, 4.65, 4.8, 0.9, AUTHOR, 11.5, WHITE)
            add_text(slide, 7.05, 4.65, 5.2, 0.9, "Научный руководитель:\n____________________", 11.5, WHITE)
            add_text(slide, 0.78, 6.62, 2.3, 0.25, CITY_YEAR, 10.5, WHITE)
            continue

        add_title(slide, data["title"])
        if "blocks" in data:
            add_blocks(slide, data["blocks"])
        elif "left" in data and "right" in data:
            add_two_columns(slide, data["left"], data["right"])
        elif "roles" in data:
            add_roles(slide, data["roles"], data["note"])
        elif "flow" in data:
            add_flow(slide, data["flow"], data["details"])
        elif "columns" in data:
            add_columns(slide, data["columns"], data["note"])
        elif "table" in data:
            add_table(slide, data["table"])
        elif "demo" in data:
            add_demo(slide, data["demo"], data["note"])
        elif "sources" in data:
            add_sources(slide, data["sources"])
        add_footer(slide, i)

    prs.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    build()
