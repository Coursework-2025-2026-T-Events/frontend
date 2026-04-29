from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
REF_DIR = Path(r"C:\Users\Office\Downloads\Telegram Desktop")
REFERENCE = REF_DIR / "Чанышева_Презентация_Т-евентс.pptx"
if not REFERENCE.exists():
    REFERENCE = max(REF_DIR.glob("*Т-евентс*.pptx"), key=lambda p: p.stat().st_size)

OUT = ROOT / "docs" / "t_events_frontend_reference_style_exact.pptx"

NAVY = RGBColor(16, 45, 105)
LIME = RGBColor(220, 255, 5)
LAV = RGBColor(223, 199, 242)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(16, 16, 16)
GRAY = RGBColor(243, 245, 248)

FOOTER_TITLE = "«Клиентская часть приложения T-Events для участия\nв образовательных активностях c элементами геймификации»"
AUTHOR = "Автор: ____________________\nОП «Программная инженерия»\nГруппа: __________"
CITY = "Москва, 2026"


def set_text(shape, text):
    if not hasattr(shape, "text"):
        return
    shape.text = text


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def add_rect(slide, x, y, w, h, fill=WHITE, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    return shape


def add_text(slide, x, y, w, h, text, size=14, color=BLACK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(slide, x, y, w, lines, size=13, gap=0.32):
    for line in lines:
        add_text(slide, x, y, w, Inches(0.28), "• " + line, size, BLACK)
        y += Inches(gap)
    return y


def cover(slide, x, y, w, h):
    add_rect(slide, x, y, w, h, WHITE)


def update_common(slide, n):
    for shape in slide.shapes:
        if not hasattr(shape, "text") or not shape.text.strip():
            continue
        text = shape.text.strip()
        if "Чанышева" in text or "БПИ237" in text:
            set_text(shape, AUTHOR)
        elif text.startswith("Москва"):
            set_text(shape, CITY)
        elif "Разработка приложения для проведения" in text or "Приложение Т-Events" in text:
            set_text(shape, FOOTER_TITLE)
        elif text.isdigit():
            set_text(shape, str(n))


def add_flow(slide, labels, x, y, w, h, font_size=10):
    step_w = w / len(labels) * 0.82
    gap = w / len(labels) * 0.18
    cur = x
    for i, label in enumerate(labels):
        fill = LAV if i % 2 == 0 else GRAY
        add_rect(slide, cur, y, step_w, h, fill, NAVY)
        add_text(slide, cur + Inches(0.03), y + Inches(0.14), step_w - Inches(0.06), Inches(0.28), label, font_size, NAVY, True, PP_ALIGN.CENTER)
        if i < len(labels) - 1:
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, cur + step_w, y + h / 2, cur + step_w + gap, y + h / 2)
            conn.line.color.rgb = NAVY
            conn.line.width = Pt(1.2)
        cur += step_w + gap


def add_module_grid(slide, x, y, modules):
    cell_w = Inches(2.2)
    cell_h = Inches(0.72)
    for i, (name, desc) in enumerate(modules):
        row = i // 3
        col = i % 3
        cx = x + col * Inches(2.55)
        cy = y + row * Inches(1.05)
        add_rect(slide, cx, cy, cell_w, cell_h, LAV if i % 2 == 0 else GRAY, NAVY)
        add_text(slide, cx + Inches(0.08), cy + Inches(0.08), cell_w - Inches(0.16), Inches(0.18), name, 10.5, NAVY, True, PP_ALIGN.CENTER)
        add_text(slide, cx + Inches(0.08), cy + Inches(0.34), cell_w - Inches(0.16), Inches(0.22), desc, 7.8, BLACK, False, PP_ALIGN.CENTER)


def add_image_or_box(slide, img_path, x, y, w, h, caption):
    add_rect(slide, x, y, w, h, WHITE, LAV)
    if img_path.exists():
        slide.shapes.add_picture(str(img_path), x, y, width=w, height=h)
    else:
        add_text(slide, x + Inches(0.1), y + Inches(0.15), w - Inches(0.2), h - Inches(0.3), caption, 12, NAVY, True, PP_ALIGN.CENTER)


def main():
    prs = Presentation(str(REFERENCE))

    # Common footer/header replacement.
    for i, slide in enumerate(prs.slides, 1):
        update_common(slide, i)

    # 1. Cover.
    s = prs.slides[0]
    set_text(s.shapes[0], "«Клиентская часть приложения T-Events для участия в образовательных активностях c элементами геймификации»")
    set_text(s.shapes[2], "Выполнил(а):\n____________________\nОП «Программная инженерия»,\nгруппа __________")
    set_text(s.shapes[4], "Курсовая работа | Программный проект | Frontend-часть")
    set_text(s.shapes[7], "«T-Events Frontend Application for Interactive Educational Activities Featuring Gamification»")

    # 2. Domain description.
    s = prs.slides[1]
    set_text(s.shapes[1], "Описание предметной области")
    set_text(
        s.shapes[2],
        "Актуальность:\n"
        "Компании проводят стендовые активности на мероприятиях: Дни карьеры, День открытых дверей, партнерские фестивали и др.\n\n"
        "Интерес участника:\n"
        "\t- быстро выбрать мероприятие и направление;\n"
        "\t- пройти игровые задания;\n"
        "\t- увидеть прогресс и получить фирменную продукцию.\n\n"
        "Интерес Т-Образования:\n"
        "\t- вовлечь аудиторию в образовательные активности;\n"
        "\t- формализовать правила получения мерча;\n"
        "\t- уменьшить ручную нагрузку на сотрудников стойки.\n\n"
        "Клиентская часть должна связать участие, прогресс, QR-награду и подтвержденную выдачу в единый пользовательский процесс.",
    )

    # 3. Goal and tasks.
    s = prs.slides[2]
    set_text(
        s.shapes[2],
        "Цель курсовой работы:\n"
        "Реализовать клиентскую часть приложения T-Events, которая обеспечивает участие в образовательных интерактивных мероприятиях, прохождение игровых активностей, получение QR-наград и управление мероприятиями.\n\n"
        "Задачи курсовой работы:\n"
        "1. Спроектировать frontend-архитектуру с учетом ролей participant, stander, admin.\n"
        "2. Реализовать пользовательские пути участника, стендиста и администратора.\n"
        "3. Интегрировать frontend с API авторизации, мероприятий, игр, наград и администрирования.\n"
        "4. Реализовать игровые сессии для question_answer и quiz.\n"
        "5. Реализовать QR/redeem-сценарий выдачи призов.\n"
        "6. Добавить runtime-проверку критичных API-контрактов.",
    )

    # 4. Analogs.
    s = prs.slides[3]
    set_text(
        s.shapes[2],
        "Проект для Т-Образования является узкоспециализированным.\n\n"
        "- Прямой аналог: T-education_bot (*Математический бот). Он закрывает часть активностей, но не дает полноценного веб-интерфейса с ролями, направлениями, QR-выдачей и администрированием.\n"
        "- Косвенные аналоги: Stepik, Quizizz, Duolingo. Они сильны в обучении и игровых заданиях, но не решают задачу физического мероприятия и выдачи фирменной продукции.\n\n"
        "Вывод: уникальность T-Events во frontend-части заключается в связке «игра -> прогресс -> право на приз -> QR-проверка -> подтвержденная выдача».",
    )

    # 5. Functional requirements analysis.
    s = prs.slides[4]
    set_text(
        s.shapes[2],
        "В начале работы над frontend-частью необходимо было определить пользовательские сценарии и требования к интерфейсам.\n\n"
        "Задачи при анализе предметной области:\n"
        "\t1) Определить пользовательский путь участника мероприятия\n"
        "\t2) Определить рабочий путь стендиста на стойке выдачи\n"
        "\t3) Определить пользовательский путь администратора\n"
        "\t4) Выявить критичные состояния: истекшая сессия, недоступная награда, повторная выдача, неготовое мероприятие\n\n"
        "Результат:\n"
        "\tБыли выделены три роли приложения, общий набор сущностей и основные сценарии: участие в играх, накопление прогресса, генерация QR-кода, проверка выдачи и управление мероприятием.\n"
        "\tГлавная «боль» frontend-части заключается в необходимости согласовать несколько бизнес-процессов в одном интерфейсе без рассинхронизации состояния.",
    )

    # 6. Participant path - keep same slide type, replace diagram area.
    s = prs.slides[5]
    set_text(s.shapes[7], "Пользовательский путь участника мероприятия (клиента):")
    set_text(s.shapes[8], "Рис. 1 Пользовательский путь участника")
    cover(s, Inches(0.62), Inches(2.32), Inches(10.95), Inches(4.0))
    add_flow(s, ["Каталог", "Мероприятие", "Направление", "Игры", "Сессия", "Прогресс", "QR"], Inches(0.85), Inches(3.0), Inches(10.4), Inches(0.58), 8.5)
    add_bullets(
        s,
        Inches(0.9),
        Inches(4.0),
        Inches(9.8),
        [
            "участник выбирает мероприятие и направление участия;",
            "frontend запускает или продолжает игровую сессию через backend;",
            "после ответа обновляются состояние сессии и прогресс направления;",
            "при достижении порога становится доступна QR-награда.",
        ],
        12.4,
        0.38,
    )

    # 7. Admin + stander path - same slide type.
    s = prs.slides[6]
    set_text(s.shapes[6], "Пользовательский путь администратора и стендиста:")
    set_text(
        s.shapes[8],
        "Основной функционал администратора:\n"
        "Работа с мероприятиями, расписанием и статусами\n"
        "Добавление направлений в мероприятие\n"
        "Назначение игр направлениям\n"
        "Настройка порогов малой и большой награды\n"
        "Проверка готовности публикации\n"
        "Публикация и архивирование мероприятий\n\n"
        "Основной функционал стендиста:\n"
        "Сканирование QR-кода или ручной redeem-код\n"
        "Preview выдачи перед подтверждением\n"
        "Подтверждение выдачи и просмотр журнала",
    )
    set_text(s.shapes[9], "Рис. 2 Пользовательский путь администратора и стендиста")
    cover(s, Inches(0.65), Inches(2.42), Inches(5.75), Inches(3.75))
    add_flow(s, ["Событие", "Настройки", "Игры", "Readiness", "Публикация"], Inches(0.75), Inches(3.0), Inches(5.3), Inches(0.52), 8.2)
    add_flow(s, ["QR/код", "Preview", "Confirm", "Inventory"], Inches(0.95), Inches(4.35), Inches(4.9), Inches(0.52), 8.5)

    # 8-10 UI slides: preserve exact screenshot-slide composition, replace pictures/captions.
    imgs = {
        "auth": ROOT / "public" / "images" / "auth-background.png",
        "catalog": ROOT / "public" / "images" / "events-catalog-hero.png",
        "choose": ROOT / "public" / "images" / "step-choose-event.png",
        "play": ROOT / "public" / "images" / "step-play-games.png",
        "prize": ROOT / "public" / "images" / "step-get-prize.png",
    }

    s = prs.slides[7]
    set_text(s.shapes[6], "Пользовательский интерфейс приложения:")
    for idx in [7, 9]:
        cover(s, s.shapes[idx].left, s.shapes[idx].top, s.shapes[idx].width, s.shapes[idx].height)
    add_image_or_box(s, imgs["auth"], s.shapes[7].left, s.shapes[7].top, s.shapes[7].width, s.shapes[7].height, "Авторизация")
    add_image_or_box(s, imgs["catalog"], s.shapes[9].left, s.shapes[9].top, s.shapes[9].width, s.shapes[9].height, "Каталог мероприятий")
    set_text(s.shapes[10], "Рис. 3\nАвторизация")
    set_text(s.shapes[8], "Рис. 4\nКаталог мероприятий")
    set_text(s.shapes[11], "Рис. 5\nГлавная страница T-Events")

    s = prs.slides[8]
    set_text(s.shapes[8], "Пользовательский интерфейс приложения:")
    for idx, img_key in [(1, "choose"), (0, "play")]:
        cover(s, s.shapes[idx].left, s.shapes[idx].top, s.shapes[idx].width, s.shapes[idx].height)
        add_image_or_box(s, imgs[img_key], s.shapes[idx].left, s.shapes[idx].top, s.shapes[idx].width, s.shapes[idx].height, img_key)
    set_text(s.shapes[10], "Рис. 6\nВыбор мероприятия")
    set_text(s.shapes[9], "Рис. 7\nВыбор направления")
    set_text(s.shapes[11], "Рис. 8, 9, 10\nИгровой сценарий участника")
    set_text(s.shapes[12], "Рис. 11\nПрогресс и награда")

    s = prs.slides[9]
    set_text(s.shapes[7], "Пользовательский интерфейс приложения:")
    for idx, img_key in [(1, "play"), (0, "prize")]:
        cover(s, s.shapes[idx].left, s.shapes[idx].top, s.shapes[idx].width, s.shapes[idx].height)
        add_image_or_box(s, imgs[img_key], s.shapes[idx].left, s.shapes[idx].top, s.shapes[idx].width, s.shapes[idx].height, img_key)
    set_text(s.shapes[8], "Рис. 12\nИгровой процесс и отправка ответа")
    set_text(s.shapes[9], "Рис. 13\nQR-награда и redeem-код")

    # 11. Technology choice.
    s = prs.slides[10]
    set_text(
        s.shapes[1],
        "Использован следующий инструментарий:\n"
        "Фронтенд:\n"
        "Фреймворк: Next.js 16 App Router\n"
        "Библиотека интерфейса: React 19\n"
        "Язык: TypeScript 5 strict mode\n"
        "Стилизация: Tailwind CSS 4\n"
        "Server state: TanStack React Query 5\n"
        "Client state: Zustand\n"
        "Валидация данных: zod\n"
        "Формы: react-hook-form\n"
        "QR: qrcode.react, @zxing/browser\n"
        "Тестирование: Node test runner, Playwright\n"
        "Сборка и запуск: Docker, npm scripts",
    )
    set_text(
        s.shapes[5],
        "Обоснование выбора технологий:\n"
        "Next.js App Router выбран из-за файловой маршрутизации и route-level состояний loading/error/not-found.\n"
        "React Query выбран для отделения server state от локального состояния интерфейса.\n"
        "Zustand используется только для локального сценария участия, чтобы не дублировать backend state.\n"
        "zod применяется для runtime-проверки ответов API, потому что TypeScript не валидирует данные во время выполнения.\n"
        "TypeScript strict mode позволяет безопасно развивать большое количество DTO и пользовательских сценариев.",
    )

    # 12. Architecture.
    s = prs.slides[11]
    set_text(s.shapes[0], "Архитектура клиентского приложения")
    set_text(
        s.shapes[1],
        "Проект содержит такие компоненты, как:\n"
        "src/app/ - маршруты Next.js App Router, page.tsx, loading/error/not-found;\n"
        "src/features/auth/ - авторизация, роли, refresh, auth lifecycle;\n"
        "src/features/events/ - мероприятия, направления, игры, игровые сессии;\n"
        "src/features/reward/ - QR-награды, redeem-коды, выдача призов;\n"
        "src/features/admin/ - административный контур, публикация, audit;\n"
        "src/components/ - UI-примитивы, layout, состояния интерфейса;\n"
        "src/lib/api/ - API-клиент, DTO, обработка ошибок;\n"
        "src/lib/routes.ts - единая карта маршрутов.",
    )
    set_text(
        s.shapes[5],
        "Архитектурной особенностью является разделение route-слоя и feature-слоя.\n\n"
        "Route-слой отвечает за подключение страницы к маршруту и базовые состояния.\n\n"
        "Feature-слой содержит предметную логику: API-методы, zod-контракты, hooks, presentation helpers и модели состояния.\n\n"
        "Это позволяет не превращать страницы в крупные файлы со смешением UI, запросов, мутаций и бизнес-правил.",
    )
    set_text(s.shapes[6], "Подробнее: см. документацию проекта `docs/frontend_work_report.md` и аудит клиентской кодовой базы.")

    # 13. Client data structure / difficulties.
    s = prs.slides[12]
    set_text(s.shapes[0], "Структура клиентских данных")
    set_text(
        s.shapes[1],
        "Клиентская модель данных включает основные сущности:\n"
        "Мероприятия\n"
        "Направления\n"
        "Игровые шаблоны\n"
        "Игры в рамках мероприятия\n"
        "Игровые сессии пользователя\n"
        "Текущий вопрос и навигация по вопросам\n"
        "Прогресс игры и направления\n"
        "Пороги малой и большой награды\n"
        "QR-награда и redeem-код\n"
        "Факт выдачи приза\n"
        "Audit logs\n\n"
        "Для критичных сущностей реализованы TypeScript DTO и zod-схемы проверки API-ответов.",
    )
    set_text(
        s.shapes[5],
        "Трудности, с которыми пришлось столкнуться:\n"
        "Frontend работает с несколькими связанными состояниями: авторизация, мероприятие, направление, игра, сессия, награда и выдача.\n\n"
        "Мероприятия, направления и игры могут изменяться администратором независимо друг от друга, но для участника должны выглядеть как единый последовательный путь.\n\n"
        "Нужно было не только отрисовать данные, но и защитить интерфейс от некорректных API-ответов, истекшей сессии, повторной выдачи и неготового мероприятия.",
    )

    # 14. Diagram slide.
    s = prs.slides[13]
    set_text(s.shapes[1], "Структура\nклиентских модулей")
    set_text(s.shapes[2], "Основные frontend-модули:")
    cover(s, Inches(3.15), Inches(1.25), Inches(8.85), Inches(5.35))
    add_module_grid(
        s,
        Inches(3.4),
        Inches(1.75),
        [
            ("Auth", "session, roles, refresh"),
            ("Events", "catalog, directions"),
            ("Game session", "question_answer, quiz"),
            ("Reward", "QR, redeem, inventory"),
            ("Admin", "settings, publish"),
            ("API client", "CSRF, retry, errors"),
            ("Contracts", "zod runtime checks"),
            ("UI", "states, primitives"),
            ("Routes", "typed route map"),
        ],
    )

    # 15. Results.
    s = prs.slides[14]
    set_text(
        s.shapes[1],
        "Могу выделить несколько результатов курсовой работы:\n"
        "1. Спроектирована frontend-архитектура, рассчитанная на три роли и несколько связанных бизнес-процессов.\n"
        "2. Реализован путь участника: каталог -> направление -> игра -> прогресс -> QR-награда.\n"
        "3. Реализован путь стендиста: QR/redeem -> preview -> confirm -> журнал выдач.\n"
        "4. Реализован административный контур: настройки мероприятия, направления, игры, readiness, публикация, архивирование и audit/export.\n"
        "5. Введена runtime-проверка критичных API-контрактов.",
    )
    set_text(
        s.shapes[5],
        "Итог - получена клиентская часть приложения, которая позволяет сформировать полноценное решение первой версии совместно с серверной частью.\n\n"
        "Главный результат frontend-разработки: интерфейс стал координатором ролей, состояний и API-контрактов, а не просто набором страниц.",
    )
    set_text(s.shapes[6], "Подробнее о frontend-части: `docs/frontend_work_report.md`, `docs/frontend_middle_2026_codebase_audit.md`.")

    # 16-17 demo slides, keep slide style.
    s = prs.slides[15]
    # cover media area but keep title/footer.
    cover(s, Inches(3.7), Inches(1.22), Inches(7.95), Inches(5.15))
    add_rect(s, Inches(3.85), Inches(1.55), Inches(7.55), Inches(4.45), GRAY, LAV)
    add_text(
        s,
        Inches(4.1),
        Inches(1.95),
        Inches(7.0),
        Inches(3.4),
        "Сценарий live-демонстрации:\n\n"
        "1. Участник выбирает мероприятие и направление\n"
        "2. Запускает игровую сессию и отправляет ответ\n"
        "3. Получает прогресс и QR-награду\n"
        "4. Стендист вводит redeem-код\n"
        "5. Preview -> подтверждение выдачи -> журнал",
        18,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )

    s = prs.slides[16]
    for idx, img_key in [(4, "play"), (5, "prize")]:
        cover(s, s.shapes[idx].left, s.shapes[idx].top, s.shapes[idx].width, s.shapes[idx].height)
        add_image_or_box(s, imgs[img_key], s.shapes[idx].left, s.shapes[idx].top, s.shapes[idx].width, s.shapes[idx].height, img_key)

    # 18. Sources.
    s = prs.slides[17]
    set_text(
        s.shapes[4],
        "Интернет источники, которые были использованы во время создания проекта:\n"
        "\t1) Next.js App Router Documentation // Official Next.js Docs\n"
        "URL: https://nextjs.org/docs/app (дата обращения: 12.05.2026).\n"
        "\t2) React Documentation // Official React Docs\n"
        "URL: https://react.dev/ (дата обращения: 12.05.2026).\n"
        "\t3) TanStack Query Documentation // Official TanStack Docs\n"
        "URL: https://tanstack.com/query (дата обращения: 12.05.2026).\n"
        "\t4) Zod Documentation // Official Zod Docs\n"
        "URL: https://zod.dev/ (дата обращения: 12.05.2026).\n"
        "\t5) ZXing Browser Documentation // GitHub repository\n"
        "URL: https://github.com/zxing-js/browser (дата обращения: 12.05.2026).\n"
        "\t6) WCAG 2.2 // W3C Recommendation\n"
        "URL: https://www.w3.org/TR/WCAG22/ (дата обращения: 12.05.2026).",
    )

    # Leave slide 19 as in reference if it is an ending/blank slide, but update common already handled.
    prs.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
