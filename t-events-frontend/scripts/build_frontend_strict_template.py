from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r"C:\Users\Office\Downloads\02_HSE_Presentation_Shablon.pptx")
OUT = ROOT / "docs" / "t_events_frontend_strict_hse_template.pptx"

NAVY = RGBColor(16, 45, 105)
LIME = RGBColor(220, 255, 5)
LAV = RGBColor(223, 199, 242)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(16, 16, 16)
GRAY = RGBColor(243, 245, 248)

AUTHOR = "Автор: ____________________\nОП «Программная инженерия»\nГруппа: __________"
SUPERVISOR = "Научный руководитель:\n____________________"
FOOTER_TITLE = "«Клиентская часть приложения T-Events для участия\nв образовательных активностях c элементами геймификации»"
CITY = "Москва, 2026"

CONTENT_SLIDES = [
    ("Описание предметной области", [
        "Компании проводят стендовые активности на Днях карьеры, днях открытых дверей и партнерских мероприятиях.",
        "Участник хочет выбрать активность, пройти задания, увидеть прогресс и получить фирменную продукцию.",
        "Организатору важно формализовать правила участия, снизить ручную нагрузку и исключить повторную выдачу призов.",
        "Клиентская часть связывает участие, прогресс, QR-награду и подтвержденную выдачу в единый процесс.",
    ], "domain"),
    ("Цель и задачи", [
        "Цель: реализовать клиентскую часть T-Events для участия в мероприятиях, прохождения игровых активностей, получения QR-наград и управления мероприятиями.",
        "Спроектировать frontend-архитектуру с учетом ролей participant, stander и admin.",
        "Реализовать пользовательские пути участника, стендиста и администратора.",
        "Интегрировать API авторизации, мероприятий, игр, наград и администрирования.",
        "Реализовать игровые сессии question_answer и quiz, QR/redeem-выдачу и runtime-проверку API-контрактов.",
    ], "tasks"),
    ("Анализ аналогов", [
        "Прямой аналог: T-education_bot. Он закрывает часть активностей, но не дает полноценного веб-интерфейса с ролями, направлениями, QR-выдачей и администрированием.",
        "Косвенные аналоги: Stepik, Quizizz, Duolingo. Они сильны в обучении и игровых заданиях, но не решают задачу физического мероприятия и выдачи мерча.",
        "Вывод: T-Events уникален связкой «игра -> прогресс -> право на приз -> QR-проверка -> подтвержденная выдача».",
    ], "analogs"),
    ("Анализ предметной области. Функциональные требования", [
        "Выделены три роли приложения: участник, стендист и администратор.",
        "Участник: каталог мероприятий, выбор направления, прохождение игр, прогресс, QR-награда.",
        "Стендист: сканирование QR или ручной redeem-код, preview выдачи, подтверждение, журнал.",
        "Администратор: создание и настройка мероприятия, направления, игры, пороги наград, публикация, audit.",
        "Главная сложность frontend-части: согласовать несколько бизнес-процессов без рассинхронизации состояния.",
    ], "roles"),
    ("Анализ предметной области. Функциональные требования", [
        "Пользовательский путь участника мероприятия:",
        "Каталог -> мероприятие -> направление -> список игр -> игровая сессия -> прогресс -> QR-награда.",
        "После ответа обновляется не только текущий вопрос, но и прогресс направления.",
        "Доступность награды определяется порогами малой и большой награды.",
    ], "participant-flow"),
    ("Анализ предметной области. Функциональные требования", [
        "Пользовательский путь администратора и стендиста:",
        "Администратор: событие -> настройки -> направления -> игры -> readiness -> публикация.",
        "Стендист: QR/redeem-код -> preview выдачи -> confirm -> журнал выданных призов.",
        "Readiness-модель помогает не опубликовать неполное мероприятие.",
    ], "admin-stander-flow"),
    ("Анализ предметной области. Функциональные требования", [
        "Пользовательский интерфейс приложения:",
        "Авторизация, каталог мероприятий и карточка мероприятия.",
        "Эти экраны формируют вход в пользовательский путь и определяют доступные действия по роли пользователя.",
    ], "ui-1"),
    ("Анализ предметной области. Функциональные требования", [
        "Пользовательский интерфейс приложения:",
        "Выбор направления, список игр и прогресс направления.",
        "Интерфейс показывает не только список активностей, но и состояние прохождения и путь к награде.",
    ], "ui-2"),
    ("Анализ предметной области. Функциональные требования", [
        "Пользовательский интерфейс приложения:",
        "Игровая сессия, QR-награда и выдача стендистом.",
        "QR-код и redeem-код являются частью бизнес-процесса подтвержденной выдачи приза.",
    ], "ui-3"),
    ("Выбор технологий для реализации", [
        "Фронтенд: Next.js 16 App Router, React 19, TypeScript 5 strict mode, Tailwind CSS 4.",
        "Server state: TanStack React Query 5. Client state: Zustand.",
        "Валидация: zod, react-hook-form. QR: qrcode.react, @zxing/browser.",
        "App Router выбран для маршрутов и route-level состояний loading/error/not-found.",
        "zod используется, потому что TypeScript не проверяет реальные backend-ответы во время выполнения.",
    ], "tech"),
    ("Архитектура системы", [
        "src/app/ - маршруты Next.js App Router, page.tsx, loading/error/not-found.",
        "src/features/auth/ - авторизация, роли, refresh, auth lifecycle.",
        "src/features/events/ - мероприятия, направления, игры, игровые сессии.",
        "src/features/reward/ - QR-награды, redeem-коды, выдача призов.",
        "src/features/admin/ - настройки мероприятия, публикация, архивирование, audit.",
        "src/lib/api/ - API-клиент, DTO, обработка ошибок и runtime-контракты.",
    ], "architecture"),
    ("Структура клиентских данных", [
        "Основные сущности: мероприятие, направление, игра, игровая сессия, текущий вопрос, прогресс, награда, QR-код, факт выдачи, audit log.",
        "Трудность: несколько связанных состояний должны выглядеть для пользователя как единый путь.",
        "Игры, направления и мероприятия могут меняться независимо.",
        "Ошибки API-структуры не должны попадать в UI как корректные данные.",
    ], "data"),
    ("Структура клиентских модулей", [
        "Auth: session, roles, refresh, logout sync.",
        "Events: catalog, directions, games, game session, leaderboard.",
        "Reward: eligibility, QR, redeem, preview, redemption, inventory.",
        "Admin: event settings, directions, games, readiness, publish, archive, audit.",
        "API client: CSRF, credentials, retry after refresh, ApiError.",
        "Contracts: zod-схемы для auth, events, reward, admin.",
    ], "modules"),
    ("Результаты проделанной работы", [
        "Спроектирована frontend-архитектура, рассчитанная на три роли и несколько связанных бизнес-процессов.",
        "Реализован путь участника: каталог -> направление -> игра -> прогресс -> QR.",
        "Реализован путь стендиста: QR/redeem -> preview -> confirm -> inventory.",
        "Реализован административный контур: настройки, readiness, публикация, архивирование, audit/export.",
        "Главный результат: frontend стал координатором ролей, состояний и API-контрактов, а не просто набором страниц.",
    ], "results"),
    ("Демонстрация работы продукта", [
        "Сценарий демонстрации:",
        "1. Участник выбирает мероприятие и направление.",
        "2. Проходит игровую сессию и отправляет ответ.",
        "3. Получает прогресс и QR-награду.",
        "4. Стендист проверяет redeem-код и подтверждает выдачу.",
        "5. Администратор видит настройки, readiness и audit.",
    ], "demo"),
    ("Демонстрация работы продукта", [
        "Ключевые экраны демонстрации:",
        "Игровой процесс участника.",
        "QR-награда и redeem-код.",
        "Preview и подтверждение выдачи стендистом.",
        "Административная настройка мероприятия.",
    ], "demo-screens"),
    ("Список источников", [
        "Next.js App Router Documentation. URL: https://nextjs.org/docs/app",
        "React Documentation. URL: https://react.dev/",
        "TanStack Query Documentation. URL: https://tanstack.com/query",
        "Zod Documentation. URL: https://zod.dev/",
        "ZXing Browser Documentation. URL: https://github.com/zxing-js/browser",
        "WCAG 2.2. URL: https://www.w3.org/TR/WCAG22/",
    ], "sources"),
]


def clear_slides(prs):
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)


def set_text(shape, value):
    shape.text = value


def sorted_shapes(slide):
    return sorted(slide.shapes, key=lambda s: (s.top, s.left))


def text_box(slide, x, y, w, h, value, size=12, color=BLACK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    return shape


def bullets_text(items):
    return "\n".join("• " + item for item in items)


def find_main_shapes(slide):
    shapes = list(slide.shapes)
    title = max(shapes, key=lambda s: s.width if s.top > Inches(1.0) and s.top < Inches(2.0) else 0)
    body_candidates = [s for s in shapes if hasattr(s, "text") and s.top > Inches(2.0) and s.left < Inches(6.5)]
    body = max(body_candidates, key=lambda s: s.width * s.height)
    image_candidates = [s for s in shapes if s.left > Inches(6.0) and s.top > Inches(1.0)]
    image = max(image_candidates, key=lambda s: s.width * s.height) if image_candidates else None
    return title, body, image


def draw_right_visual(slide, image_shape, kind):
    if image_shape is None:
        return
    x, y, w, h = image_shape.left, image_shape.top, image_shape.width, image_shape.height
    rect(slide, x, y, w, h, GRAY, LAV)
    if "flow" in kind:
        labels = ["Событие", "Игра", "Прогресс", "QR", "Выдача"]
        step_w = int(w / len(labels) * 0.78)
        gap = int(w / len(labels) * 0.22)
        cur = x + Inches(0.18)
        cy = y + int(h * 0.43)
        for index, label in enumerate(labels):
            rect(slide, cur, cy, step_w, Inches(0.38), LAV if index % 2 == 0 else WHITE, NAVY)
            text_box(slide, cur + Inches(0.02), cy + Inches(0.1), step_w - Inches(0.04), Inches(0.12), label, 7.5, NAVY, True, PP_ALIGN.CENTER)
            cur += step_w + gap
    elif kind.startswith("ui") or kind.startswith("demo"):
        card_w = int(w * 0.72)
        card_x = x + int(w * 0.14)
        card_y = y + int(h * 0.18)
        rect(slide, card_x, card_y, card_w, int(h * 0.64), WHITE, LAV)
        rect(slide, card_x + Inches(0.12), card_y + Inches(0.12), card_w - Inches(0.24), Inches(0.28), NAVY)
        rect(slide, card_x + Inches(0.16), card_y + Inches(0.75), card_w - Inches(0.32), Inches(0.34), GRAY, LAV)
        rect(slide, card_x + Inches(0.16), card_y + Inches(1.25), card_w - Inches(0.32), Inches(0.34), GRAY, LAV)
        text_box(slide, card_x + Inches(0.18), card_y + int(h * 0.48), card_w - Inches(0.36), Inches(0.3), "UI screen", 13, NAVY, True, PP_ALIGN.CENTER)
    else:
        text_box(slide, x + Inches(0.25), y + int(h * 0.38), w - Inches(0.5), Inches(0.6), "T-Events\nfrontend", 22, NAVY, True, PP_ALIGN.CENTER)


def fill_cover(slide):
    shapes = list(slide.shapes)
    text_shapes = [s for s in shapes if hasattr(s, "text")]
    # Use positions from the original cover template.
    by_area = sorted(text_shapes, key=lambda s: s.width * s.height, reverse=True)
    set_text(by_area[0], "«Клиентская часть приложения T-Events для участия в образовательных активностях c элементами геймификации»")
    for s in text_shapes:
        if s is by_area[0]:
            continue
        if s.left < Inches(4.0) and s.top < Inches(2.0):
            set_text(s, "Факультет\nКомпьютерных Наук")
        elif s.left > Inches(5.5) and s.left < Inches(8.5):
            set_text(s, "Выполнил(а):\n____________________\nОП «Программная инженерия»,\nгруппа __________")
        elif s.left > Inches(8.5):
            set_text(s, SUPERVISOR)
        elif s.top > Inches(4.0):
            set_text(s, "Курсовая работа | Программный проект | Frontend-часть")


def fill_content(slide, title_value, items, kind, number):
    title_shape, body_shape, image_shape = find_main_shapes(slide)
    set_text(title_shape, title_value)
    set_text(body_shape, bullets_text(items))
    for shape in slide.shapes:
        if not hasattr(shape, "text"):
            continue
        if shape is title_shape or shape is body_shape:
            continue
        current = shape.text.strip()
        if current == "":
            if shape.top < Inches(1.0) and shape.left < Inches(3.5):
                set_text(shape, AUTHOR)
            elif shape.top < Inches(1.0) and shape.left > Inches(3.0) and shape.left < Inches(6.2):
                set_text(shape, FOOTER_TITLE)
            elif shape.top < Inches(1.0) and shape.left > Inches(6.0) and shape.left < Inches(9.0):
                set_text(shape, CITY)
        elif current.isdigit() or shape.left > Inches(10.0):
            set_text(shape, str(number))

    draw_right_visual(slide, image_shape, kind)


def build():
    prs = Presentation(str(TEMPLATE))
    cover_layout = prs.slide_layouts[0]
    text_layout = prs.slide_layouts[1]
    blank_layout = prs.slide_layouts[-1]
    clear_slides(prs)

    slide = prs.slides.add_slide(cover_layout)
    fill_cover(slide)

    for number, (title_value, items, kind) in enumerate(CONTENT_SLIDES, start=2):
        slide = prs.slides.add_slide(text_layout)
        fill_content(slide, title_value, items, kind, number)

    prs.slides.add_slide(blank_layout)
    prs.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    build()
