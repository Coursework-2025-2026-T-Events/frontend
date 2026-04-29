from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r"C:\Users\Office\Downloads\02_HSE_Presentation_Shablon.pptx")
FALLBACK_TEMPLATE = ROOT / "docs" / "hse_fcs_template.pptx"
OUT = ROOT / "docs" / "t_events_defense_presentation_hse_template.pptx"

NAVY = RGBColor(16, 45, 105)
LIME = RGBColor(220, 255, 5)
LAVENDER = RGBColor(223, 199, 242)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(16, 16, 16)
GRAY = RGBColor(243, 245, 248)


SLIDES = [
    {
        "type": "title",
        "title": "T-Events",
        "subtitle": "Клиентская часть веб-приложения для интерактивного сопровождения мероприятий",
        "footer": "Защита проекта, 10 минут",
    },
    {
        "title": "1. Проблема и необходимость",
        "bullets": [
            "На мероприятиях игровые активности, прогресс участников и выдача призов часто ведутся вручную или в разрозненных инструментах.",
            "При росте числа участников появляются риски: неверный подсчет прогресса, повторная выдача призов, ошибки сотрудников стойки.",
            "Нужен единый клиентский контур, который сопровождает путь от участия в активности до подтвержденной выдачи награды.",
        ],
        "accent": "Не каталог мероприятий, а цифровой процесс участия и выдачи призов.",
    },
    {
        "title": "2. Идея решения",
        "bullets": [
            "T-Events объединяет три роли: участник, стендист и администратор.",
            "Участник проходит игры и открывает награду; стендист проверяет QR/redeem-код; администратор настраивает мероприятие.",
            "Frontend выступает координатором связанных сценариев, состояний и API-контрактов.",
        ],
        "diagram": ["Участник", "Игра", "Прогресс", "QR", "Стендист", "Выдача"],
    },
    {
        "title": "3. Уникальность проекта",
        "bullets": [
            "Объединены игровые механики, прогресс по направлениям, право на реальную награду и административное управление событием.",
            "QR-код является частью бизнес-процесса, а не декоративным элементом интерфейса.",
            "Система поддерживает полный цикл: настройка мероприятия -> прохождение заданий -> получение QR -> проверка -> подтверждение выдачи.",
        ],
        "accent": "Frontend решает задачу согласования ролей и бизнес-состояний, а не только отображения экранов.",
    },
    {
        "title": "4. Архитектурная схема клиента",
        "columns": [
            (
                "src/app",
                [
                    "Маршруты Next.js App Router",
                    "page/loading/error/not-found",
                    "Клиентские контейнеры сценариев",
                ],
            ),
            (
                "src/features",
                [
                    "auth, events, reward, admin",
                    "API-методы и zod-контракты",
                    "Предметная логика и presentation helpers",
                ],
            ),
            (
                "src/lib + components",
                [
                    "Единый API-клиент",
                    "Маршруты, query keys, tokenStore",
                    "UI-примитивы и layout",
                ],
            ),
        ],
        "accent": "Страницы связывают сценарий, feature-слой хранит бизнес-правила.",
    },
    {
        "title": "5. Сложность: игровая сессия",
        "bullets": [
            "Сессия имеет статус, текущий вопрос, навигацию, прогресс, историю ответов и результат последнего действия.",
            "Поддержаны два движка: question_answer с текстовым ответом и quiz с вариантами ответа.",
            "После ответа нужно обновить не только игру, но и прогресс направления и доступность награды.",
        ],
        "solution": "Решение: discriminated union по engine, нормализация session DTO, вынос расчетов и presentation-логики в feature/events.",
    },
    {
        "title": "6. Сложность: награды и QR-выдача",
        "bullets": [
            "Награда зависит от прогресса по направлению: малая и большая награда открываются по разным порогам.",
            "QR содержит signed token; дополнительно показывается ручной redeem-код на случай недоступной камеры.",
            "Выдача построена как двухшаговый процесс: preview -> confirm, чтобы снизить риск ошибки сотрудника.",
        ],
        "solution": "Решение: отдельный reward-модуль, QR/redeem utilities, API-контракты для eligibility, QR, preview, redemption и inventory.",
    },
    {
        "title": "7. Сложность: авторизация и безопасность сценариев",
        "bullets": [
            "Есть разные уровни доступа: participant, stander, admin.",
            "Access token хранится в памяти, refresh выполняется через cookie, после 401 запрос повторяется автоматически.",
            "Logout/login синхронизируются между вкладками, а refresh-запросы коалесцируются, чтобы избежать гонок.",
        ],
        "solution": "Решение: централизованный API-клиент, tokenStore, AuthProvider, BroadcastChannel/storage events, route guards.",
    },
    {
        "title": "8. Сложность: надежность API-контрактов",
        "bullets": [
            "TypeScript не проверяет фактические данные, которые пришли с backend во время выполнения.",
            "Ошибка в ответе игровой сессии, награды или публикации мероприятия может нарушить бизнес-процесс.",
            "Критичные ответы валидируются runtime-схемами и классифицируются как contract_mismatch.",
        ],
        "solution": "Решение: authContracts, eventContracts, rewardContracts, adminContracts на базе zod + единая модель ApiError.",
    },
    {
        "title": "9. Административный контур",
        "bullets": [
            "Администратор управляет мероприятием: детали, расписание, часовой пояс, пороги наград, направления и игры.",
            "Публикация требует готовности разделов: расписание, награды, направления, игры.",
            "Интерфейс показывает проблемы готовности и направляет администратора к нужному разделу.",
        ],
        "solution": "Решение: settings tabs, readiness model, publish checklist, admin event/game config helpers и audit/export сценарии.",
    },
    {
        "title": "10. Демонстрация за 2 минуты",
        "bullets": [
            "Участник: каталог -> мероприятие -> направление -> игра -> ответ -> прогресс -> QR-награда.",
            "Стендист: ручной redeem-код или QR -> preview выдачи -> подтверждение -> журнал выдач.",
            "Администратор: настройки мероприятия -> направления/игры -> проверка готовности -> публикация/audit.",
        ],
        "accent": "Демонстрация показывает не отдельные страницы, а сквозной бизнес-процесс.",
    },
    {
        "title": "11. Итоги и развитие",
        "bullets": [
            "Разработана клиентская часть, которая связывает участие, игровую механику, прогресс, QR-награду и выдачу приза.",
            "Главный результат: frontend стал координатором ролей, состояний и API-контрактов в сложном пользовательском процессе.",
            "Дальнейшее развитие: новые игровые механики, склад призов, аналитика, расширенное e2e/a11y, observability и performance budgets.",
        ],
        "accent": "Практическая значимость: меньше ручной нагрузки, меньше ошибок выдачи, прозрачнее управление мероприятием.",
    },
]


def clear_existing_slides(prs: Presentation) -> None:
    slide_id_list = prs.slides._sldIdLst
    for slide_id in list(slide_id_list):
        prs.part.drop_rel(slide_id.rId)
        slide_id_list.remove(slide_id)


def get_blank_layout(prs: Presentation):
    for layout in prs.slide_layouts:
        if len(layout.placeholders) == 0:
            return layout
    return prs.slide_layouts[-1]


def add_rect(slide, x, y, w, h, color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    return shape


def add_text(slide, x, y, w, h, value, size=16, color=BLACK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = Inches(0.06)
    frame.margin_right = Inches(0.06)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = Inches(0.04)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def decorate(slide, slide_number):
    add_rect(slide, 0, 0, 13.333, 0.18, LIME)
    add_rect(slide, 0, 7.22, 13.333, 0.28, NAVY)
    add_text(slide, 12.3, 7.17, 0.6, 0.25, str(slide_number), 10, WHITE, True, PP_ALIGN.RIGHT)


def add_bullets(slide, bullets):
    y = 1.45
    for bullet in bullets:
        add_text(slide, 0.78, y, 11.0, 0.52, "• " + bullet, 16.5, BLACK)
        y += 0.74


def add_note_box(slide, note):
    if note.startswith("Решение:"):
        add_rect(slide, 0.65, 5.78, 11.95, 0.7, NAVY)
        add_text(slide, 0.82, 5.88, 11.55, 0.38, note, 13.5, WHITE)
    else:
        add_rect(slide, 0.65, 5.78, 11.95, 0.7, LAVENDER)
        add_text(slide, 0.82, 5.9, 11.55, 0.34, note, 14.5, NAVY, True)


def add_regular_slide(prs, layout, slide_data, slide_number):
    slide = prs.slides.add_slide(layout)
    decorate(slide, slide_number)
    add_text(slide, 0.55, 0.48, 11.7, 0.55, slide_data["title"], 27, NAVY, True)

    if "columns" in slide_data:
        for column_index, (heading, items) in enumerate(slide_data["columns"]):
            x = 0.65 + column_index * 4.15
            add_rect(slide, x, 1.42, 3.65, 3.75, GRAY, LAVENDER)
            add_text(slide, x + 0.12, 1.55, 3.42, 0.35, heading, 18, NAVY, True)
            y = 2.1
            for item in items:
                add_text(slide, x + 0.18, y, 3.28, 0.48, "• " + item, 13.5, BLACK)
                y += 0.74
    else:
        add_bullets(slide, slide_data["bullets"])

    if "diagram" in slide_data:
        x = 0.65
        for index, label in enumerate(slide_data["diagram"]):
            add_rect(slide, x, 4.65, 1.33, 0.52, LAVENDER if index % 2 == 0 else GRAY, NAVY)
            add_text(slide, x + 0.03, 4.73, 1.27, 0.2, label, 11.5, NAVY, True, PP_ALIGN.CENTER)
            if index < len(slide_data["diagram"]) - 1:
                add_text(slide, x + 1.33, 4.72, 0.34, 0.2, "->", 14, NAVY, True, PP_ALIGN.CENTER)
            x += 1.66

    note = slide_data.get("solution") or slide_data.get("accent")
    if note:
        add_note_box(slide, note)


def add_title_slide(prs, layout, slide_data, slide_number):
    slide = prs.slides.add_slide(layout)
    add_rect(slide, 0, 0, 13.333, 7.5, NAVY)
    add_rect(slide, 0, 0, 13.333, 0.26, LIME)
    add_rect(slide, 8.85, 0.58, 2.8, 0.75, LAVENDER)
    add_text(slide, 0.75, 1.38, 6.8, 0.9, slide_data["title"], 58, LIME, True)
    add_text(slide, 0.75, 2.35, 8.8, 1.0, slide_data["subtitle"], 27, WHITE)
    add_text(slide, 0.75, 5.72, 5.0, 0.35, slide_data["footer"], 15, WHITE)
    add_rect(slide, 0, 7.22, 13.333, 0.28, NAVY)
    add_text(slide, 12.3, 7.17, 0.6, 0.25, str(slide_number), 10, WHITE, True, PP_ALIGN.RIGHT)


def main():
    template_path = TEMPLATE if TEMPLATE.exists() else FALLBACK_TEMPLATE
    prs = Presentation(str(template_path)) if template_path.exists() else Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    clear_existing_slides(prs)
    layout = get_blank_layout(prs)

    for index, slide_data in enumerate(SLIDES, start=1):
        if slide_data.get("type") == "title":
            add_title_slide(prs, layout, slide_data, index)
        else:
            add_regular_slide(prs, layout, slide_data, index)

    prs.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
