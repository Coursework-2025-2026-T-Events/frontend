from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r"C:\Users\Office\Downloads\02_HSE_Presentation_Shablon.pptx")
FALLBACK = ROOT / "docs" / "hse_fcs_template.pptx"
OUT = ROOT / "docs" / "t_events_frontend_logical_hse_clean.pptx"

NAVY = RGBColor(16, 45, 105)
LIME = RGBColor(220, 255, 5)
LAV = RGBColor(223, 199, 242)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(16, 16, 16)
GRAY = RGBColor(243, 245, 248)

AUTHOR = "Автор: ____________________\nОП «Программная инженерия»\nГруппа: __________"
SUPERVISOR = "Научный руководитель:\n____________________"
TITLE = "«Клиентская часть приложения T-Events для участия\nв образовательных активностях c элементами геймификации»"
CITY = "Москва, 2026"


SLIDES = [
    ("cover", "Клиентская часть приложения T-Events", "Frontend-часть программного проекта"),
    ("text_image", "Описание предметной области", [
        "Компании проводят стендовые активности на Днях карьеры, днях открытых дверей и партнерских мероприятиях.",
        "Участник хочет быстро выбрать активность, пройти задания, увидеть прогресс и получить фирменную продукцию.",
        "Организатору важно формализовать правила участия, снизить ручную нагрузку и исключить повторную выдачу призов.",
        "Frontend-часть должна связать участие, прогресс, QR-награду и выдачу в единый пользовательский процесс.",
    ]),
    ("text_image", "Цель и задачи", [
        "Цель: реализовать клиентскую часть T-Events для участия в мероприятиях, прохождения игровых активностей, получения QR-наград и управления мероприятиями.",
        "Задачи: спроектировать frontend-архитектуру по ролям participant, stander, admin.",
        "Реализовать пользовательские пути участника, стендиста и администратора.",
        "Интегрировать API авторизации, мероприятий, игр, наград и администрирования.",
        "Реализовать игровые сессии question_answer и quiz, QR/redeem-выдачу и runtime-проверку API-контрактов.",
    ]),
    ("analogs", "Анализ аналогов", [
        "Прямой аналог: T-education_bot. Закрывает часть активностей, но не дает полноценного веб-интерфейса с ролями, направлениями, QR-выдачей и администрированием.",
        "Косвенные аналоги: Stepik, Quizizz, Duolingo. Они сильны в обучении и игровых заданиях, но не решают задачу физического мероприятия и выдачи мерча.",
        "Вывод: T-Events уникален связкой «игра -> прогресс -> право на приз -> QR-проверка -> подтвержденная выдача».",
    ]),
    ("text_image", "Анализ предметной области. Функциональные требования", [
        "На этапе анализа были выделены три пользовательские роли: участник, стендист и администратор.",
        "Участник: каталог мероприятий, выбор направления, прохождение игр, прогресс, QR-награда.",
        "Стендист: сканирование QR или ручной redeem-код, preview выдачи, подтверждение, журнал.",
        "Администратор: создание и настройка мероприятия, направления, игры, пороги наград, публикация, audit.",
        "Главная сложность frontend-части: согласовать несколько бизнес-процессов без рассинхронизации состояния.",
    ]),
    ("flow", "Анализ предметной области. Функциональные требования", "Пользовательский путь участника мероприятия", ["Каталог", "Мероприятие", "Направление", "Игры", "Сессия", "Прогресс", "QR"]),
    ("flow_text", "Анализ предметной области. Функциональные требования", "Пользовательский путь администратора и стендиста", ["Событие", "Настройки", "Игры", "Readiness", "Публикация"], [
        "Администратор управляет мероприятием как составным объектом: детали, расписание, направления, игры, пороги наград.",
        "Стендист работает с операционным сценарием: QR/redeem -> preview -> confirm -> inventory.",
        "Readiness-модель не позволяет опубликовать неполное мероприятие.",
    ]),
    ("ui", "Анализ предметной области. Функциональные требования", "Пользовательский интерфейс приложения", ["Авторизация", "Каталог мероприятий", "Карточка мероприятия"]),
    ("ui", "Анализ предметной области. Функциональные требования", "Пользовательский интерфейс приложения", ["Выбор направления", "Список игр", "Прогресс направления"]),
    ("ui", "Анализ предметной области. Функциональные требования", "Пользовательский интерфейс приложения", ["Игровая сессия", "QR-награда", "Выдача стендистом"]),
    ("two_col", "Выбор технологий для реализации", [
        ("Использованный инструментарий", [
            "Next.js 16 App Router",
            "React 19, TypeScript 5 strict",
            "Tailwind CSS 4",
            "TanStack React Query 5",
            "Zustand",
            "zod, react-hook-form",
            "qrcode.react, @zxing/browser",
            "Playwright, Docker",
        ]),
        ("Обоснование выбора", [
            "App Router удобен для маршрутов и route-level состояний.",
            "React Query отделяет server state от локального UI state.",
            "Zustand используется только для локального сценария участия.",
            "zod проверяет реальные backend-ответы во время выполнения.",
            "TypeScript strict снижает риск ошибок в DTO и сценариях.",
        ]),
    ]),
    ("two_col", "Архитектура системы", [
        ("Компоненты проекта", [
            "src/app/ - маршруты, page, loading/error/not-found",
            "src/features/auth/ - авторизация и роли",
            "src/features/events/ - мероприятия, игры, сессии",
            "src/features/reward/ - QR, redeem, выдача",
            "src/features/admin/ - настройки, публикация, audit",
            "src/lib/api/ - API-клиент и DTO",
        ]),
        ("Архитектурная особенность", [
            "Route-слой отвечает за подключение сценария к URL.",
            "Feature-слой хранит предметную логику, API-методы, contracts, hooks и presentation helpers.",
            "Такое разделение не позволяет страницам превращаться в файлы со смешением UI, запросов и бизнес-правил.",
        ]),
    ]),
    ("two_col", "Структура клиентских данных", [
        ("Основные сущности", [
            "Мероприятие, направление, игра",
            "Игровая сессия, текущий вопрос, навигация",
            "Прогресс игры и направления",
            "Пороги малой и большой награды",
            "QR-награда, redeem-код",
            "Факт выдачи, audit logs",
        ]),
        ("Трудности", [
            "Несколько связанных состояний должны выглядеть для пользователя как единый путь.",
            "Игры, направления и мероприятия могут меняться независимо.",
            "Ошибки API-структуры не должны попадать в UI как корректные данные.",
            "Нужно учитывать истекшую сессию, неготовое мероприятие и повторную выдачу.",
        ]),
    ]),
    ("diagram", "Структура клиентских модулей", ["Auth", "Events", "Game session", "Reward", "Admin", "API client", "Contracts", "UI", "Routes"]),
    ("two_col", "Результаты проделанной работы", [
        ("Результаты", [
            "Спроектирована frontend-архитектура по ролям и сценариям.",
            "Реализован путь участника: каталог -> игра -> прогресс -> QR.",
            "Реализован путь стендиста: QR/redeem -> preview -> confirm -> inventory.",
            "Реализован административный контур: настройки, readiness, публикация, audit/export.",
        ]),
        ("Итог", [
            "Клиентская часть формирует полноценное решение первой версии совместно с серверной частью.",
            "Главный результат: frontend стал координатором ролей, состояний и API-контрактов, а не просто набором страниц.",
        ]),
    ]),
    ("demo", "Демонстрация работы продукта", [
        "Участник выбирает мероприятие и направление",
        "Проходит игровую сессию и отправляет ответ",
        "Получает прогресс и QR-награду",
        "Стендист проверяет redeem-код и подтверждает выдачу",
        "Администратор видит настройки, readiness и audit",
    ]),
    ("ui", "Демонстрация работы продукта", "Ключевые экраны демонстрации", ["Игровой процесс", "QR-награда", "Административная настройка"]),
    ("sources", "Список источников", [
        "Next.js App Router Documentation. URL: https://nextjs.org/docs/app",
        "React Documentation. URL: https://react.dev/",
        "TanStack Query Documentation. URL: https://tanstack.com/query",
        "Zod Documentation. URL: https://zod.dev/",
        "ZXing Browser Documentation. URL: https://github.com/zxing-js/browser",
        "WCAG 2.2. URL: https://www.w3.org/TR/WCAG22/",
    ]),
    ("thanks", "", ""),
]


def clear_slides(prs):
    ids = prs.slides._sldIdLst
    for slide_id in list(ids):
        prs.part.drop_rel(slide_id.rId)
        ids.remove(slide_id)


def blank_layout(prs):
    for layout in prs.slide_layouts:
        if len(layout.placeholders) == 0:
            return layout
    return prs.slide_layouts[-1]


def rect(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    return shape


def text(slide, x, y, w, h, value, size=13, color=BLACK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    p = frame.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def footer(slide, n):
    text(slide, 1.25, 0.55, 2.2, 0.5, AUTHOR, 6.8, NAVY)
    text(slide, 6.85, 0.86, 1.8, 0.2, CITY, 7, NAVY)
    text(slide, 3.95, 0.58, 2.8, 0.48, TITLE, 7, NAVY)
    text(slide, 11.25, 0.63, 0.6, 0.35, str(n), 13, NAVY, True, PP_ALIGN.CENTER)


def title(slide, value):
    text(slide, 0.65, 1.58, 9.4, 0.65, value, 24, NAVY, True)


def bullets(slide, x, y, w, items, size=13, gap=0.43):
    for item in items:
        text(slide, x, y, w, 0.35, "• " + item, size, BLACK)
        y += gap


def flow(slide, labels, x, y, w, h):
    step_w = w / len(labels) * 0.82
    gap = w / len(labels) * 0.18
    cur = x
    for i, label in enumerate(labels):
        rect(slide, cur, y, step_w, h, LAV if i % 2 == 0 else GRAY, NAVY)
        text(slide, cur + 0.02, y + 0.14, step_w - 0.04, 0.22, label, 9, NAVY, True, PP_ALIGN.CENTER)
        if i < len(labels) - 1:
            connector = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(cur + step_w),
                Inches(y + h / 2),
                Inches(cur + step_w + gap),
                Inches(y + h / 2),
            )
            connector.line.color.rgb = NAVY
            connector.line.width = Pt(1)
        cur += step_w + gap


def draw_mock(slide, x, y, w, h, caption):
    rect(slide, x, y, w, h, WHITE, LAV)
    rect(slide, x + 0.12, y + 0.14, w - 0.24, 0.34, NAVY)
    rect(slide, x + 0.18, y + 0.72, w - 0.36, 0.42, GRAY, LAV)
    rect(slide, x + 0.18, y + 1.28, w - 0.36, 0.42, GRAY, LAV)
    text(slide, x + 0.15, y + h - 0.48, w - 0.3, 0.25, caption, 10.5, NAVY, True, PP_ALIGN.CENTER)


def build():
    template = TEMPLATE if TEMPLATE.exists() else FALLBACK
    prs = Presentation(str(template)) if template.exists() else Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    clear_slides(prs)
    layout = blank_layout(prs)

    for i, data in enumerate(SLIDES, 1):
        kind = data[0]
        slide = prs.slides.add_slide(layout)
        if kind == "cover":
            rect(slide, 0, 0, 13.333, 7.5, WHITE)
            rect(slide, 0, 0, 13.333, 0.25, LIME)
            text(slide, 1.25, 1.0, 3.2, 0.45, "Факультет\nКомпьютерных Наук", 16, NAVY, True)
            text(slide, 1.25, 2.35, 9.6, 1.25, f"«{data[1]} для участия в образовательных активностях c элементами геймификации»", 27, NAVY, True)
            text(slide, 1.25, 4.75, 4.8, 0.8, "Курсовая работа | Программный проект | Frontend-часть", 14, NAVY)
            text(slide, 6.4, 1.05, 2.6, 0.85, AUTHOR, 10.5, NAVY)
            text(slide, 9.25, 1.05, 2.8, 0.85, SUPERVISOR, 10.5, NAVY)
            text(slide, 9.25, 6.1, 2.0, 0.3, CITY, 12, NAVY)
            text(slide, 1.25, 5.18, 7.5, 0.35, "«T-Events Frontend Application for Interactive Educational Activities Featuring Gamification»", 11, NAVY)
            continue

        footer(slide, i)
        title(slide, data[1])
        if kind in {"text_image", "analogs"}:
            bullets(slide, 0.75, 2.35, 7.4, data[2], 12.6, 0.48)
            rect(slide, 8.6, 2.15, 3.3, 3.2, GRAY, LAV)
            text(slide, 8.85, 3.3, 2.8, 0.6, "T-Events\nfrontend", 21, NAVY, True, PP_ALIGN.CENTER)
        elif kind == "flow":
            text(slide, 0.75, 2.3, 7.5, 0.35, data[2] + ":", 15, BLACK)
            flow(slide, data[3], 0.9, 3.15, 10.9, 0.58)
            bullets(slide, 1.0, 4.25, 9.9, [
                "выбор мероприятия и направления сохраняется в локальном сценарии участия;",
                "игровая сессия запускается или продолжается через backend;",
                "после ответа обновляется прогресс игры и направления;",
                "QR-награда становится доступной после достижения порогов.",
            ], 12.4, 0.42)
        elif kind == "flow_text":
            text(slide, 0.75, 2.3, 7.8, 0.35, data[2] + ":", 15, BLACK)
            flow(slide, data[3], 0.9, 3.0, 5.7, 0.55)
            flow(slide, ["QR/код", "Preview", "Confirm", "Inventory"], 0.9, 4.15, 5.7, 0.55)
            bullets(slide, 7.25, 2.55, 4.6, data[4], 12.2, 0.55)
        elif kind == "ui":
            text(slide, 0.75, 2.25, 7.5, 0.35, data[2] + ":", 15, BLACK)
            x = 1.05
            for cap in data[3]:
                draw_mock(slide, x, 2.8, 3.25, 2.65, cap)
                x += 3.75
        elif kind == "two_col":
            left, right = data[2]
            text(slide, 0.75, 2.35, 4.9, 0.35, left[0] + ":", 15, BLACK, True)
            bullets(slide, 0.85, 2.85, 5.2, left[1], 11.8, 0.38)
            text(slide, 6.65, 2.35, 4.9, 0.35, right[0] + ":", 15, BLACK, True)
            bullets(slide, 6.75, 2.85, 5.0, right[1], 11.8, 0.45)
        elif kind == "diagram":
            text(slide, 0.75, 2.3, 5.8, 0.35, "Основные frontend-модули:", 15, BLACK)
            labels = data[2]
            x0, y0 = 3.05, 2.25
            for idx, label in enumerate(labels):
                x = x0 + (idx % 3) * 2.65
                y = y0 + (idx // 3) * 1.05
                rect(slide, x, y, 2.15, 0.7, LAV if idx % 2 == 0 else GRAY, NAVY)
                text(slide, x + 0.08, y + 0.22, 2.0, 0.2, label, 10.5, NAVY, True, PP_ALIGN.CENTER)
        elif kind == "demo":
            text(slide, 3.7, 1.7, 7.5, 0.55, "Демонстрация работы\nпродукта", 27, NAVY, True, PP_ALIGN.CENTER)
            bullets(slide, 3.95, 3.0, 6.9, data[2], 14, 0.48)
        elif kind == "sources":
            bullets(slide, 0.75, 2.35, 10.8, data[2], 12.2, 0.48)
        elif kind == "thanks":
            text(slide, 3.0, 3.2, 7.3, 0.6, "Спасибо за внимание", 32, NAVY, True, PP_ALIGN.CENTER)

    prs.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    build()
